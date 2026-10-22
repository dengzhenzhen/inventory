import uuid
import pytest
from conftest import login


def post(
    c,
    kind="purchase",
    quantity=10,
    price=10,
    date="2026-01-01",
    code="A",
    size="M",
    key=None,
    **extra,
):
    line = {
        "style_code": code,
        "size": size,
        "quantity": quantity,
        "price": price,
        **extra,
    }
    r = c.post(
        "/api/documents",
        headers={"Idempotency-Key": key or uuid.uuid4().hex},
        json={"kind": kind, "date": date, "remark": "测试原因", "lines": [line]},
    )
    assert r.status_code == 200, r.text
    return r.json()["data"]


def pieces(c):
    return c.get("/api/stats/summary").json()["data"]["pieces"]


def test_session_revocation_and_initial_password(client):
    data = login(client)
    assert data["must_change_password"]
    assert client.get("/api/stock").status_code == 403
    old = client.headers["Authorization"]
    assert (
        client.put(
            "/api/auth/password",
            json={"old_password": "admin123", "new_password": "newpass123"},
        ).status_code
        == 200
    )
    assert client.get("/api/auth/me").status_code == 401
    login(client, "newpass123")
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").status_code == 401
    client.headers["Authorization"] = old
    assert client.get("/api/auth/me").status_code == 401


def test_rename_preserves_stock(auth):
    d = post(auth)
    p = d["lines"][0]["product_id"]
    s = d["lines"][0]["size_id"]
    assert auth.put(f"/api/products/{p}", json={"style_code": "B"}).status_code == 200
    assert auth.put(f"/api/dict/size/{s}", json={"name": "中码"}).status_code == 200
    r = auth.get(
        "/api/stock/current", params={"style_code": "B", "size": "中码"}
    ).json()["data"]
    assert r["stock"] == 10
    detail = auth.get(f"/api/documents/{d['id']}").json()["data"]
    assert detail["lines"][0]["style_snapshot"] == "A"


def test_historical_cost_and_explicit_recompute(auth):
    post(auth)
    sale = post(auth, "sale", 2, 30, "2026-01-02")
    before = auth.get("/api/stats/summary", params={"date_to": "2026-01-31"}).json()[
        "data"
    ]
    post(auth, "purchase", 10, 20, "2026-02-01")
    after = auth.get("/api/stats/summary", params={"date_to": "2026-01-31"}).json()[
        "data"
    ]
    assert before["profit"] == after["profit"] == 40
    post(auth, "purchase", 10, 20, "2025-12-31")
    preview = auth.get("/api/documents/cost-preview").json()["data"]
    assert preview["changes"]
    assert (
        auth.post(
            "/api/documents/cost-recompute",
            json={"fingerprint": preview["fingerprint"]},
        ).status_code
        == 200
    )
    detail = auth.get(f"/api/documents/{sale['id']}").json()["data"]
    assert detail["lines"][0]["cost"] == 30


def test_preview_stale(auth):
    post(auth)
    p = auth.get("/api/documents/cost-preview").json()["data"]
    post(auth, code="B")
    assert (
        auth.post(
            "/api/documents/cost-recompute", json={"fingerprint": p["fingerprint"]}
        ).status_code
        == 409
    )


def test_idempotent_and_atomic(auth):
    first = post(auth, key="same")
    second = post(auth, key="same")
    assert first["id"] == second["id"]
    assert pieces(auth) == 10
    r = auth.post(
        "/api/documents",
        headers={"Idempotency-Key": "same"},
        json={
            "kind": "purchase",
            "date": "2026-01-01",
            "lines": [{"style_code": "B", "quantity": 1}],
        },
    )
    assert r.status_code == 409
    r = auth.post(
        "/api/documents",
        headers={"Idempotency-Key": "fail"},
        json={
            "kind": "sale",
            "date": "2026-01-02",
            "lines": [
                {"style_code": "A", "size": "M", "quantity": 1},
                {"style_code": "absent", "size": "M", "quantity": 1},
            ],
        },
    )
    assert r.status_code == 404
    assert pieces(auth) == 10


def test_returns_and_void(auth):
    purchase = post(auth)
    sale = post(auth, "sale", 3, 30, "2026-01-02")
    r = post(
        auth, "sale_return", 2, 0, "2026-01-03", original_line_id=sale["lines"][0]["id"]
    )
    assert pieces(auth) == 9
    assert r["lines"][0]["cost"] == 20
    assert r["lines"][0]["unit_price"] == 30
    too_many = auth.post(
        "/api/documents",
        headers={"Idempotency-Key": "excess"},
        json={
            "kind": "sale_return",
            "date": "2026-01-04",
            "lines": [
                {
                    "style_code": "A",
                    "size": "M",
                    "quantity": 2,
                    "original_line_id": sale["lines"][0]["id"],
                }
            ],
        },
    )
    assert too_many.status_code == 409
    assert (
        auth.post(
            f"/api/documents/{sale['id']}/void", json={"reason": "更正"}
        ).status_code
        == 409
    )
    assert (
        auth.post(f"/api/documents/{r['id']}/void", json={"reason": "更正"}).status_code
        == 200
    )
    assert (
        auth.post(
            f"/api/documents/{sale['id']}/void", json={"reason": "更正"}
        ).status_code
        == 200
    )
    assert pieces(auth) == 10
    assert (
        auth.post(
            f"/api/documents/{sale['id']}/void", json={"reason": "重复"}
        ).status_code
        == 200
    )
    assert pieces(auth) == 10
    post(
        auth,
        "purchase_return",
        1,
        0,
        "2026-01-05",
        original_line_id=purchase["lines"][0]["id"],
    )
    assert pieces(auth) == 9


def test_count_conflict(auth):
    post(auth)
    v = auth.get("/api/stock/current", params={"style_code": "A", "size": "M"}).json()[
        "data"
    ]["stock_version"]
    count = post(auth, "count", 1, 0, "2026-01-03", stock_version=v, counted=7)
    assert pieces(auth) == 7
    r = auth.post(
        "/api/documents",
        headers={"Idempotency-Key": "count"},
        json={
            "kind": "count",
            "date": "2026-01-04",
            "remark": "盘点",
            "lines": [
                {
                    "style_code": "A",
                    "size": "M",
                    "quantity": 1,
                    "counted": 5,
                    "stock_version": v,
                }
            ],
        },
    )
    assert r.status_code == 409


def test_negative_and_missing_cost(auth):
    post(auth, "purchase", 1, 10)
    sale = post(auth, "sale", 3, 30, "2026-01-02")
    assert sale["lines"][0]["cost_status"] == "provisional"
    assert pieces(auth) == -2
    post(auth, "purchase", 5, 20, "2026-01-03")
    assert (
        auth.get(f"/api/documents/{sale['id']}").json()["data"]["lines"][0][
            "cost_status"
        ]
        == "provisional"
    )
    preview = auth.get("/api/documents/cost-preview").json()["data"]
    assert (
        auth.post(
            "/api/documents/cost-recompute",
            json={"fingerprint": preview["fingerprint"]},
        ).status_code
        == 200
    )
    settled = auth.get(f"/api/documents/{sale['id']}").json()["data"]["lines"][0]
    assert settled["cost_status"] == "known"
    assert settled["cost"] == 50
    returned = post(
        auth, "sale_return", 1, 0, "2026-01-04", original_line_id=sale["lines"][0]["id"]
    )
    assert returned["lines"][0]["cost"] == 16.67
    assert auth.post("/api/products", json={"style_code": "B"}).status_code == 200
    missing = post(auth, "sale", 1, 30, "2026-01-04", code="B")
    assert missing["lines"][0]["cost"] is None
    assert auth.get("/api/stats/summary").json()["data"]["profit"] is None


@pytest.mark.parametrize(
    "patch", [{"style_code": " "}, {"date": "not-date"}, {"price": -1}]
)
def test_validation(auth, patch):
    data = {
        "kind": "purchase",
        "date": patch.get("date", "2026-01-01"),
        "lines": [
            {
                "style_code": patch.get("style_code", "A"),
                "size": "M",
                "quantity": 1,
                "price": patch.get("price", 1),
            }
        ],
    }
    assert (
        auth.post(
            "/api/documents", headers={"Idempotency-Key": uuid.uuid4().hex}, json=data
        ).status_code
        == 422
    )


def test_pagination_and_export(auth):
    from app.database import SessionLocal
    from app.models import Product

    with SessionLocal() as db:
        db.add_all([Product(style_code=f"S{i}", retail_cents=100) for i in range(505)])
        db.commit()
    assert auth.get("/api/products", params={"page_size": 500}).status_code == 422
    assert (
        auth.get("/api/products", params={"keyword": "S504"}).json()["data"]["total"]
        == 1
    )
    assert auth.get("/api/products/by-code/S504").status_code == 200
    post(auth, code="=1+1")
    post(auth, code="OTHER")
    r = auth.get("/api/export/purchases.xlsx", params={"style_code": "=1+1"})
    from openpyxl import load_workbook
    from io import BytesIO

    rows = list(load_workbook(BytesIO(r.content)).active.rows)
    assert len(rows) == 2
    assert rows[1][3].value == "=1+1"
    assert rows[1][3].data_type == "s"


def test_concurrent_idempotency(auth):
    from concurrent.futures import ThreadPoolExecutor

    def submit(_):
        return post(auth, key="concurrent-key")["id"]

    with ThreadPoolExecutor(max_workers=2) as pool:
        ids = list(pool.map(submit, range(2)))
    assert ids[0] == ids[1]
    assert pieces(auth) == 10


def test_decimal_amounts_and_missing_inventory_cost(auth):
    post(auth, quantity=3, price=0.1)
    sale = post(auth, "sale", 3, 0.3, "2026-01-02")
    assert sale["lines"][0]["cost"] == 0.3
    result = auth.get("/api/stats/summary").json()["data"]
    assert result["sale_amount"] == 0.9
    assert result["profit"] == 0.6


def test_multi_filter_document_export(auth):
    post(auth, code="A")
    other = post(auth, code="B")
    auth.post(f"/api/documents/{other['id']}/void", json={"reason": "取消"})
    params = {"kind": "purchase", "status": "void", "style_code": "B"}
    listed = auth.get("/api/documents", params=params).json()["data"]
    assert listed["total"] == 1
    from openpyxl import load_workbook
    from io import BytesIO

    exported = auth.get("/api/export/documents.xlsx", params=params)
    rows = list(load_workbook(BytesIO(exported.content)).active.rows)
    assert len(rows) == 2
    assert rows[1][2].value == "void"
    assert rows[1][3].value == "B"
